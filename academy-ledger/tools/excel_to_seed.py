"""수기 엑셀 장부(총괄 + 월별 입출금 시트)를 장부 앱 데이터(JSON)로 변환한다.

사용: python excel_to_seed.py 장부.xlsx out_dir/
출력: months_YYYY-MM.json (월별 거래), rules.json (거래처→항목 규칙), students.json
"""
import sys, json, re, datetime, hashlib, collections, os
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)
tx = []

def norm_dt(v):
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d') if v.time() == datetime.time() else v.strftime('%Y-%m-%d %H:%M:%S')
    m = re.match(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})\s*(\d{1,2}:\d{2}(?::\d{2})?)?', str(v).strip())
    if not m:
        return None
    y, mo, d, t = m.groups()
    out = f'{y}-{int(mo):02d}-{int(d):02d}'
    return f'{out} {t if t.count(":") == 2 else t + ":00"}' if t else out

def add(dt, direction, amount, name, cat, memo='', bal=None):
    if not amount or not dt:
        return
    try:
        amount = int(round(float(amount)))
    except (TypeError, ValueError):
        return
    if amount <= 0:
        return
    name = (str(name).strip() if name else '') or '(이름없음)'
    tx.append(dict(dt=dt, dir=direction, amount=amount, name=name, memo=memo or '',
                   bal=int(bal) if isinstance(bal, (int, float)) else None, cat=cat or '', src='excel'))

# 2~3월 출금: 일자, 월, 출금(이름), 금액, 비고(항목)
ws = wb['2월~3월 출금']
for r in ws.iter_rows(min_row=2, max_col=5, values_only=True):
    d, _, name, amt, cat = r
    if d:
        add(norm_dt(d), 'out', amt, name, cat)

# 2~3월 입금: 입금일시, 입금자명, 금액
ws = wb['2월~3월 입금']
for r in ws.iter_rows(min_row=2, max_col=3, values_only=True):
    d, name, amt = r
    if d:
        add(norm_dt(d), 'in', amt, name, '수강료')

# 4월 이후: B 거래일자, D 출금, E 입금, F 잔액, G 거래내용, H 거래기록사항, I 항목
for sn in wb.sheetnames:
    if not sn.endswith('입출금'):
        continue
    ws = wb[sn]
    for r in ws.iter_rows(min_row=2, min_col=2, max_col=9, values_only=True):
        d, _, out_amt, in_amt, bal, memo, name, cat = r
        if not d:
            continue
        dt = norm_dt(d)
        cat = (cat or '').strip() if isinstance(cat, str) else ''
        if out_amt:
            add(dt, 'out', out_amt, name, cat or '기타', memo, bal)
        if in_amt:
            add(dt, 'in', in_amt, name, '통장이동' if cat == '통장이동' else '수강료', memo, bal)

# id: 같은 날 같은 금액/이름 거래가 여러 건일 수 있어 순번을 붙인다
seen = collections.Counter()
for t in tx:
    base = f"{t['dt']}|{t['dir']}|{t['amount']}|{t['name']}"
    seen[base] += 1
    t['id'] = hashlib.sha1(f"{base}|{seen[base]}".encode()).hexdigest()[:12]

months = collections.defaultdict(list)
for t in tx:
    months[t['dt'][:7]].append(t)

os.makedirs(out, exist_ok=True)
for m, items in sorted(months.items()):
    items.sort(key=lambda t: t['dt'])
    json.dump({'month': m, 'items': items}, open(f'{out}/months_{m}.json', 'w'), ensure_ascii=False)

# 규칙: 거래처별로 가장 많이 쓰인 항목 (방향별)
cnt = collections.defaultdict(collections.Counter)
for t in tx:
    if t['cat']:
        cnt[(t['dir'], t['name'].replace(' ', ''))][t['cat']] += 1
rules = {'in': {}, 'out': {}}
for (d, n), c in cnt.items():
    cat = c.most_common(1)[0][0]
    if d == 'in' and cat == '수강료':
        continue  # 입금 기본값이라 규칙 불필요
    rules[d][n] = cat
json.dump(rules, open(f'{out}/rules.json', 'w'), ensure_ascii=False, indent=1)

# 재원생수 (총괄 6행)
ws = wb['총괄']
students = {}
for col in range(2, 14):
    label, v = ws.cell(2, col).value, ws.cell(6, col).value
    if isinstance(v, (int, float)) and label:
        mnum = int(re.sub(r'\D', '', label))
        students[f'2025-{mnum:02d}'] = int(v)
json.dump(students, open(f'{out}/students.json', 'w'), ensure_ascii=False)

for m, items in sorted(months.items()):
    inc = sum(t['amount'] for t in items if t['dir'] == 'in' and t['cat'] != '통장이동')
    exp = sum(t['amount'] for t in items if t['dir'] == 'out' and t['cat'] != '통장이동')
    print(m, len(items), '수입', inc, '지출', exp)
print('rules', len(rules['in']), len(rules['out']), 'students', students)
